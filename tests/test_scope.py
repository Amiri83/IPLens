"""Per-account scope: validation, resolution and its effect on every view and export.

Placeholder data only (10.0.x.x, account 123456789012, example names).
"""

import io
import json

import pytest
from openpyxl import load_workbook

from iplens import queries
from iplens import scope as scope_mod
from iplens import suggestions as sugg_mod
from iplens.db import closing
from iplens.web import create_app

VPC_A, VPC_B = "vpc-0example0000a", "vpc-0example0000b"
SUB_A1, SUB_A2, SUB_B1 = "subnet-000000a1", "subnet-000000a2", "subnet-000000b1"
ENI_A1, ENI_A2, ENI_B1 = "eni-00000000a1", "eni-00000000a2", "eni-00000000b1"


def _snapshot(builder, db_path, account_ref=1):
    b = builder(db_path, account_ref=account_ref)
    b.vpc(VPC_A, "10.0.0.0/17", name="example-vpc-a")
    b.vpc(VPC_B, "10.0.128.0/17", name="example-vpc-b")
    b.subnet(SUB_A1, VPC_A, "10.0.1.0/24", name="example-private-a")
    b.subnet(SUB_A2, VPC_A, "10.0.2.0/24", az="us-east-1b", name="example-public-a")
    b.subnet(SUB_B1, VPC_B, "10.0.129.0/24", name="example-private-b")
    b.eni(ENI_A1, SUB_A1, ["10.0.1.10", "10.0.1.200"], status="available", description="a1")
    b.eni(ENI_A2, SUB_A2, ["10.0.2.10"], status="available", description="a2")
    b.eni(ENI_B1, SUB_B1, ["10.0.129.10"], status="available", description="b1")
    return b


# -- unit -------------------------------------------------------------------------------


def test_validate_normalises_and_rejects_bad_input():
    cfg = scope_mod.validate(
        {
            "vpcs": [VPC_A, VPC_A],
            "subnets": [SUB_A1],
            "subnet_patterns": "example-*\n\n  *-b ",
            "cidr_mode": "exclude",
            "cidrs": "10.0.1.7/24, 10.0.2.0/25",
        }
    )
    assert cfg.vpcs == [VPC_A]
    assert cfg.subnet_patterns == ["example-*", "*-b"]
    assert cfg.cidrs == ["10.0.1.0/24", "10.0.2.0/25"]
    assert cfg.count == 6
    assert scope_mod.Scope.from_json(cfg.to_json()) == cfg
    for bad, msg in (
        ({"vpcs": ["not-a-vpc"]}, "invalid VPC"),
        ({"subnets": ["subnet-X;drop"]}, "invalid subnet"),
        ({"cidrs": "10.0.0.0/33"}, "invalid CIDR"),
        ({"cidrs": "fd00::/8"}, "IPv4"),
        ({"vpc_mode": "maybe"}, "include or exclude"),
    ):
        with pytest.raises(ValueError, match=msg):
            scope_mod.validate(bad)


def test_resolve_dimensions(db_path, snapshot_builder):
    snap = _snapshot(snapshot_builder, db_path)

    def resolved(**doc):
        with closing(db_path) as conn:
            return scope_mod.resolve(conn, snap.id, scope_mod.validate(doc))

    assert resolved() is scope_mod.UNSCOPED and not resolved().active
    r = resolved(vpcs=[VPC_A])
    assert r.subnet_ids == {SUB_A1, SUB_A2} and r.vpc_ids == {VPC_A}
    r = resolved(vpc_mode="exclude", vpcs=[VPC_A])
    assert r.subnet_ids == {SUB_B1} and r.vpc_ids == {VPC_B}
    r = resolved(subnet_patterns="EXAMPLE-private-*")  # case-insensitive wildcard
    assert r.subnet_ids == {SUB_A1, SUB_B1} and r.vpc_ids == {VPC_A, VPC_B}
    r = resolved(subnet_mode="exclude", subnets=[SUB_A1], subnet_patterns="*-b")
    assert r.subnet_ids == {SUB_A2} and r.vpc_ids == {VPC_A}
    r = resolved(cidrs="10.0.1.0/25")
    assert r.subnet_ids == {SUB_A1} and r.ip_ok("10.0.1.10") and not r.ip_ok("10.0.1.200")
    r = resolved(cidr_mode="exclude", cidrs="10.0.128.0/17\n10.0.1.128/25")
    assert r.subnet_ids == {SUB_A1, SUB_A2}  # a partly excluded subnet stays
    assert r.ip_ok("10.0.1.10") and not r.ip_ok("10.0.1.200")


def test_queries_and_suggestions_respect_scope(db_path, snapshot_builder):
    snap = _snapshot(snapshot_builder, db_path)
    with closing(db_path) as conn:
        r = scope_mod.resolve(
            conn,
            snap.id,
            scope_mod.validate({"vpcs": [VPC_A], "cidrs": "10.0.1.0/25\n10.0.2.0/24"}),
        )
        assert [v.vpc_id for v in queries.vpc_tree(conn, snap.id, r)] == [VPC_A]
        ips = [row["ip"] for row in queries.ip_list(conn, snap.id, scope=r)]
        assert ips == ["10.0.1.10", "10.0.2.10"]
        assert queries.owner_breakdown(conn, snap.id, r) == {"ec2": 2}
        ctx = sugg_mod.build_context(conn, snap.id, r)
        assert set(ctx.enis) == {ENI_A1, ENI_A2}
        assert ctx.eni_ip_count[ENI_A1] == 1  # 10.0.1.200 is out of range
        assert set(ctx.subnets) == {SUB_A1, SUB_A2}
        assert ctx.all_subnet_ids == {SUB_A1, SUB_A2, SUB_B1}
        data = queries.visual_data(conn, snap.id, VPC_B, scope=r)
        assert data is None  # VPC B is out of scope


# -- web --------------------------------------------------------------------------------


@pytest.fixture
def app(home):
    return create_app(home, testing=True)


@pytest.fixture
def client(app):
    c = app.test_client()
    c.get("/")
    return c


def _post(client, url, data=None, **kw):
    with client.session_transaction() as s:
        token = s["csrf"]
    return client.post(url, data={**(data or {}), "csrf_token": token}, **kw)


def _db(app):
    return app.extensions["iplens"]["paths"].db_path


def _pages(client) -> dict[str, str]:
    xlsx = load_workbook(io.BytesIO(client.get("/ips/export.xlsx").data))
    return {
        "overview": client.get("/").data.decode(),
        "ips": client.get("/ips").data.decode(),
        "xlsx": " ".join(str(c.value) for row in xlsx.active.iter_rows() for c in row),
        "visual": client.get("/visual").data.decode(),
        "visual_data": client.get("/visual/data.json").data.decode(),
        "suggestions": client.get("/suggestions").data.decode(),
    }


def test_scope_page_applies_everywhere_and_clears(app, client, snapshot_builder):
    _snapshot(snapshot_builder, _db(app))
    page = client.get("/scope").data.decode()
    assert f'<option value="{VPC_A}" >' in page and SUB_B1 in page
    assert "scope:" not in page.split("<main>", 1)[0]

    resp = _post(
        client,
        "/scope",
        {"vpc_mode": "exclude", "vpcs": [VPC_B], "subnet_patterns": "example-private-*"},
    )
    assert resp.status_code == 302
    page = client.get("/scope").data.decode()
    assert "Active scope: 2 filters" in page
    assert "1 of 2 VPCs and 1 of 3 subnets in scope" in page
    assert f'<option value="{VPC_B}" selected>' in page

    for name, body in _pages(client).items():
        assert VPC_B not in body and "10.0.129.10" not in body, name
        assert "10.0.2.10" not in body and ENI_A2 not in body, name
        assert "scope: 2 filters" in body or name in ("xlsx", "visual_data"), name
    assert "10.0.1.10" in _pages(client)["ips"]
    assert client.get(f"/subnets/{SUB_A2}").status_code == 404
    assert "outside the active scope" in client.get(f"/enis/{ENI_B1}").data.decode()
    assert client.get(f"/subnets/{SUB_A1}").status_code == 200

    resp = _post(client, "/scope/clear", {"next": "/ips"})
    assert resp.headers["Location"] == "/ips"
    pages = _pages(client)
    assert "10.0.129.10" in pages["ips"] and "10.0.2.10" in pages["xlsx"]
    assert "scope:" not in pages["ips"].split("<main>", 1)[0]


def test_scope_is_per_account(app, client, snapshot_builder):
    resp = _post(
        client,
        "/accounts/new",
        {"display_name": "example-b", "region": "us-east-1", "auth_mode": "env"},
    )
    assert resp.status_code == 302
    _snapshot(snapshot_builder, _db(app), account_ref=1)
    _snapshot(snapshot_builder, _db(app), account_ref=2)
    _post(client, "/scope", {"vpcs": [VPC_A]})
    assert "scope: 1 filter" in client.get("/").data.decode()
    _post(client, "/accounts/active", {"account_id": "2", "next": "/"})
    page = client.get("/ips").data.decode()
    assert "scope:" not in page.split("<main>", 1)[0] and "10.0.129.10" in page
    with closing(_db(app)) as conn:
        stored = {
            r["account_ref"]: json.loads(r["config"]) for r in conn.execute("SELECT * FROM scopes")
        }
    assert list(stored) == [1] and stored[1]["vpcs"] == [VPC_A]


def test_invalid_scope_is_rejected(app, client, snapshot_builder):
    _snapshot(snapshot_builder, _db(app))
    page = _post(client, "/scope", {"cidrs": "10.0.0.0/99"}, follow_redirects=True).data.decode()
    assert "Scope not saved: invalid CIDR" in page
    assert "scope:" not in page.split("<main>", 1)[0]
    assert _post(client, "/scope", {"cidrs": "10.0.0.0/16"}).status_code == 302
    assert _post(client, "/scope/clear", {"next": "//evil.example/"}).headers["Location"] == "/"
