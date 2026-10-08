"""Multi-account web flows: credentials UI, secret hygiene, scoping, history, Visual state
and exports. Placeholder data only (10.0.x.x, account 123456789012, fake credentials)."""

import io
import json
from datetime import UTC, datetime, timedelta

import boto3
import pytest
from botocore.exceptions import ClientError
from moto import mock_aws
from openpyxl import load_workbook

from iplens.db import closing
from iplens.web import create_app

FAKE_KEY_ID = "AKIAEXAMPLEEXAMPLE00"
FAKE_TEMP_KEY_ID = "ASIAEXAMPLEEXAMPLE00"
FAKE_SECRET = "example/secret/value/for/tests/only/0000"
FAKE_SESSION_TOKEN = "FakeSessionTokenForTestsOnly0000000000000000000000Example"
SECRETS = (FAKE_SECRET, FAKE_SESSION_TOKEN, FAKE_KEY_ID, FAKE_TEMP_KEY_ID)

VPC_A, SUBNET_A, ENI_A = "vpc-0example000000a", "subnet-0000000a", "eni-000000000a"
VPC_B, SUBNET_B, ENI_B = "vpc-0example000000b", "subnet-0000000b", "eni-000000000b"


@pytest.fixture
def app(home):
    return create_app(home, testing=True)


@pytest.fixture
def client(app):
    c = app.test_client()
    c.get("/")
    return c


def _csrf(client) -> str:
    with client.session_transaction() as s:
        return s["csrf"]


def _post(client, url, data=None, **kw):
    data = dict(data or {})
    data["csrf_token"] = _csrf(client)
    return client.post(url, data=data, **kw)


def _db(app):
    return app.extensions["iplens"]["paths"].db_path


def _log_text(app) -> str:
    return (app.extensions["iplens"]["log_dir"] / "iplens.log").read_text()


def _sts_json(expires_at: datetime) -> str:
    return json.dumps(
        {
            "Credentials": {
                "AccessKeyId": FAKE_TEMP_KEY_ID,
                "SecretAccessKey": FAKE_SECRET,
                "SessionToken": FAKE_SESSION_TOKEN,
                "Expiration": expires_at.isoformat(timespec="seconds"),
            }
        },
        indent=4,
    )


def _add_account(client, **fields) -> int:
    data = {"display_name": "", "region": "us-east-1", "auth_mode": "env", **fields}
    resp = _post(client, "/accounts/new", data)
    assert resp.status_code == 302, resp.data
    app = client.application
    return max(a.id for a in app.extensions["iplens"]["accounts"].list())


def _activate(client, account_id: int) -> None:
    resp = _post(client, "/accounts/active", {"account_id": str(account_id), "next": "/"})
    assert resp.status_code == 302


def _assert_no_secrets(text: str) -> None:
    for secret in SECRETS:
        assert secret not in text, f"secret material leaked: {secret[:4]}…"


# -- credentials UI -------------------------------------------------------------------


@pytest.mark.parametrize(
    "paste",
    [
        f"{FAKE_TEMP_KEY_ID} {FAKE_SECRET} {FAKE_SESSION_TOKEN}",
        f"export AWS_ACCESS_KEY_ID={FAKE_TEMP_KEY_ID}\n"
        f"export AWS_SECRET_ACCESS_KEY={FAKE_SECRET}\n"
        f"export AWS_SESSION_TOKEN={FAKE_SESSION_TOKEN}\n",
        "sts-json",  # built at run time: the expiry is relative to now
    ],
    ids=["three-fields", "export-block", "sts-json"],
)
def test_temporary_credentials_all_paste_formats(app, client, paste):
    if paste == "sts-json":
        paste = _sts_json(datetime.now(UTC) + timedelta(minutes=42, seconds=30))
    acct_id = _add_account(
        client, display_name="example-temp", auth_mode="temporary", temporary_paste=paste
    )
    acct = app.extensions["iplens"]["accounts"].get(acct_id, with_secret=True)
    assert (acct.access_key_id, acct.secret_access_key, acct.session_token) == (
        FAKE_TEMP_KEY_ID,
        FAKE_SECRET,
        FAKE_SESSION_TOKEN,
    )
    page = client.get("/settings").data.decode()
    _assert_no_secrets(page)
    if paste.startswith("{"):
        assert "expires in 42m" in page
    else:
        assert "expiry unknown" in page


def test_expiry_shown_in_settings_form_and_header(app, client):
    acct_id = _add_account(
        client,
        auth_mode="temporary",
        temporary_paste=_sts_json(datetime.now(UTC) + timedelta(minutes=30, seconds=30)),
    )
    _activate(client, acct_id)
    for url in ("/", "/settings", f"/accounts/{acct_id}/edit"):
        page = client.get(url).data.decode()
        assert "expires in 30m" in page, url
        _assert_no_secrets(page)


def test_secrets_absent_from_every_page_and_the_log(app, client, snapshot_builder):
    keys_id = _add_account(
        client,
        display_name="example-keys",
        auth_mode="keys",
        access_key_id=FAKE_KEY_ID,
        secret_access_key=FAKE_SECRET,
    )
    temp_id = _add_account(
        client,
        display_name="example-temp",
        auth_mode="temporary",
        temporary_paste=_sts_json(datetime.now(UTC) + timedelta(hours=1)),
        memory_only="on",
    )
    snapshot_builder(_db(app), account_ref=temp_id).vpc(VPC_A, "10.0.0.0/16").subnet(
        SUBNET_A, VPC_A, "10.0.1.0/24"
    )
    _activate(client, temp_id)
    # a failing refresh and test connection, plus a rejected form submission
    app.extensions["iplens"]["gateway_factory"] = _raise_expired
    _post(client, "/refresh")
    _post(client, f"/accounts/{keys_id}/test")
    bad = _post(
        client,
        "/accounts/new",
        {
            "auth_mode": "temporary",
            "region": "us-east-1",
            "temporary_paste": f"{FAKE_TEMP_KEY_ID} {FAKE_SECRET}",  # 2 fields: rejected
        },
    )
    assert bad.status_code == 400
    _assert_no_secrets(bad.data.decode())

    urls = [
        "/",
        "/settings",
        f"/accounts/{keys_id}/edit",
        f"/accounts/{temp_id}/edit",
        "/accounts/new",
        "/ips",
        "/visual",
        "/visual/data.json",
        "/rules",
        "/rules/new",
        "/suggestions",
        "/logs",
        "/logs?level=DEBUG",
    ]
    for url in urls:
        resp = client.get(url)
        assert resp.status_code == 200, url
        _assert_no_secrets(resp.data.decode())

    log_text = _log_text(app)
    _assert_no_secrets(log_text)
    assert "account created" in log_text and "refresh failed" in log_text


def test_masked_key_id_and_blank_keeps_secret(app, client):
    acct_id = _add_account(
        client, auth_mode="keys", access_key_id=FAKE_KEY_ID, secret_access_key=FAKE_SECRET
    )
    page = client.get(f"/accounts/{acct_id}/edit").data.decode()
    assert "AKIA************LE00" in page and "stored, leave blank to keep" in page
    _post(
        client,
        f"/accounts/{acct_id}/edit",
        {"auth_mode": "keys", "region": "eu-west-1", "display_name": "example-renamed"},
    )
    acct = app.extensions["iplens"]["accounts"].get(acct_id, with_secret=True)
    assert (acct.access_key_id, acct.secret_access_key, acct.region) == (
        FAKE_KEY_ID,
        FAKE_SECRET,
        "eu-west-1",
    )


def test_account_form_validation_rerenders_without_secrets(client):
    resp = _post(
        client,
        "/accounts/new",
        {
            "display_name": "example-keep-me",
            "auth_mode": "keys",
            "region": "us-east-1",
            "secret_access_key": FAKE_SECRET,
        },
    )
    assert resp.status_code == 400
    page = resp.data.decode()
    assert "access key id and secret are required" in page
    assert 'value="example-keep-me"' in page  # non-secret fields survive
    _assert_no_secrets(page)
    resp = _post(client, "/accounts/new", {"auth_mode": "profile", "region": "us-east-1"})
    assert b"choose a named profile" in resp.data


def test_env_mode_note_and_profile_dropdown(client, tmp_path, monkeypatch):
    cfg = tmp_path / "aws-config"
    cfg.write_text(
        "[profile example-readonly]\nregion = us-east-2\n\n"
        "[profile example-sso]\nsso_session = example-org\n\n"
        "[sso-session example-org]\nsso_start_url = https://example.invalid/start\n"
    )
    monkeypatch.setenv("AWS_CONFIG_FILE", str(cfg))
    page = client.get("/accounts/new").data.decode()
    assert (
        "Uses the <b>process environment</b>" in page and "No\n      credentials are stored" in page
    )
    assert '<option value="example-readonly" >example-readonly</option>' in page
    assert '<option value="example-sso" >example-sso (SSO)</option>' in page
    assert "example-org" not in page.split('name="profile"')[1].split("</select>")[0]
    acct_id = _add_account(client, auth_mode="profile", profile="example-sso")
    page = client.get(f"/accounts/{acct_id}/edit").data.decode()
    assert '<option value="example-sso" selected>example-sso (SSO)</option>' in page


def test_memory_only_survives_requests_not_restarts(app, home, client):
    acct_id = _add_account(
        client,
        auth_mode="temporary",
        temporary_paste=f"{FAKE_TEMP_KEY_ID} {FAKE_SECRET} {FAKE_SESSION_TOKEN}",
        memory_only="on",
    )
    with closing(_db(app)) as conn:
        dump = "\n".join(conn.iterdump())
    assert FAKE_SECRET not in dump and FAKE_SESSION_TOKEN not in dump
    assert "memory only" in client.get("/settings").data.decode()

    restarted = create_app(home, testing=True).test_client()
    restarted.get("/")
    _activate(restarted, acct_id)
    page = _post(restarted, "/refresh", follow_redirects=True).data.decode()
    assert "paste them again" in page
    assert f'<a href="/accounts/{acct_id}/edit">edit account</a>' in page


# -- expired credentials at collection time ---------------------------------------------------


def _raise_expired(_account):
    raise ClientError(
        {"Error": {"Code": "ExpiredToken", "Message": "The security token is expired"}},
        "DescribeVpcs",
    )


@pytest.mark.parametrize("code", ["ExpiredToken", "InvalidClientTokenId"])
def test_refresh_with_rejected_token_asks_for_new_credentials(app, client, code):
    acct_id = _add_account(
        client,
        auth_mode="temporary",
        temporary_paste=_sts_json(datetime.now(UTC) + timedelta(hours=1)),
    )
    _activate(client, acct_id)

    def rejected(_account):
        raise ClientError({"Error": {"Code": code, "Message": "example"}}, "DescribeVpcs")

    app.extensions["iplens"]["gateway_factory"] = rejected
    page = _post(client, "/refresh", follow_redirects=True).data.decode()
    assert "Refresh failed: credentials expired, paste new ones" in page
    assert f'<a href="/accounts/{acct_id}/edit">edit account</a>' in page

    page = _post(client, f"/accounts/{acct_id}/test", follow_redirects=True).data.decode()
    assert "Connection failed: credentials expired, paste new ones" in page
    assert f'<a href="/accounts/{acct_id}/edit">edit account</a>' in page


def test_refresh_with_locally_expired_credentials_never_calls_aws(app, client):
    acct_id = _add_account(
        client,
        auth_mode="temporary",
        temporary_paste=_sts_json(datetime.now(UTC) - timedelta(minutes=5)),
    )
    _activate(client, acct_id)
    page = _post(client, "/refresh", follow_redirects=True).data.decode()
    assert "Refresh failed: credentials expired, paste new ones" in page
    # the header and the account list flag it too
    assert f'<a class="ctx-warn" href="/accounts/{acct_id}/edit">credentials expired' in page
    assert "credentials expired, paste new ones" in client.get("/settings").data.decode()

    # pasting fresh credentials clears the problem
    _post(
        client,
        f"/accounts/{acct_id}/edit",
        {
            "auth_mode": "temporary",
            "region": "us-east-1",
            "temporary_paste": _sts_json(datetime.now(UTC) + timedelta(minutes=50, seconds=30)),
        },
    )
    assert "expires in 50m" in client.get("/").data.decode()


# -- account scoping ------------------------------------------------------------------


@pytest.fixture
def two_accounts(app, client, snapshot_builder, ips):
    """Account 1 (migrated default) and a second account, each with its own snapshot."""
    b_id = _add_account(client, display_name="example-b", region="eu-west-1")
    a = snapshot_builder(_db(app), account_ref=1)
    a.vpc(VPC_A, "10.0.0.0/16", name="example-vpc-a").subnet(SUBNET_A, VPC_A, "10.0.1.0/24")
    a.eni(ENI_A, SUBNET_A, ips("10.0.1.0", 10, 2), status="available", description="detached a")
    b = snapshot_builder(_db(app), "eu-west-1", account_ref=b_id, account_id="123456789012")
    b.vpc(VPC_B, "10.0.0.0/16", name="example-vpc-b").subnet(SUBNET_B, VPC_B, "10.0.2.0/24")
    b.eni(ENI_B, SUBNET_B, ips("10.0.2.0", 20, 3), status="available", description="detached b")
    _activate(client, 1)
    return {"a": a, "b": b, "b_id": b_id}


def _scoped_pages(client) -> dict[str, str]:
    xlsx = load_workbook(io.BytesIO(client.get("/ips/export.xlsx").data))
    return {
        "overview": client.get("/").data.decode(),
        "ips": client.get("/ips").data.decode(),
        "xlsx": " ".join(str(c.value) for row in xlsx.active.iter_rows() for c in row),
        "visual": client.get("/visual").data.decode(),
        "visual_data": client.get("/visual/data.json").data.decode(),
        "suggestions": client.get("/suggestions").data.decode(),
    }


def test_views_are_scoped_to_the_active_account(client, two_accounts):
    pages = _scoped_pages(client)
    for name, page in pages.items():
        assert VPC_A in page or "10.0.1.10" in page or ENI_A in page, name
        assert VPC_B not in page and "10.0.2.20" not in page and ENI_B not in page, name
    assert "10.0.1.10" in pages["ips"] and "10.0.1.10" in pages["xlsx"]
    assert ENI_A in pages["suggestions"]
    assert client.get(f"/subnets/{SUBNET_B}").status_code == 404
    assert client.get(f"/enis/{ENI_B}").status_code == 404

    _activate(client, two_accounts["b_id"])
    pages = _scoped_pages(client)
    for name, page in pages.items():
        assert VPC_A not in page and "10.0.1.10" not in page and ENI_A not in page, name
    assert "10.0.2.20" in pages["ips"] and ENI_B in pages["suggestions"]
    assert json.loads(pages["visual_data"])["vpc"]["vpc_id"] == VPC_B
    assert "<b>example-b (123456789012)</b> · eu-west-1" in pages["overview"]
    assert client.get(f"/enis/{ENI_B}").status_code == 200
    assert client.get(f"/enis/{ENI_A}").status_code == 404


def test_account_selector_in_top_bar(client, two_accounts):
    page = client.get("/ips?q=10.0").data.decode()
    assert 'action="/accounts/active"' in page
    assert '<option value="1" selected>123456789012 · us-east-1</option>' in page
    assert f'<option value="{two_accounts["b_id"]}" >example-b · eu-west-1</option>' in page
    assert 'name="next" value="/ips?q=10.0"' in page
    # detail pages send you back to Discovery after switching
    assert 'name="next" value="/"' in client.get(f"/subnets/{SUBNET_A}").data.decode()


@pytest.mark.parametrize("target", ["//evil.example/", "https://evil.example/", "/\\evil.example"])
def test_account_switch_rejects_open_redirect(client, two_accounts, target):
    resp = _post(client, "/accounts/active", {"account_id": "1", "next": target})
    assert resp.headers["Location"] == "/"


def test_account_switch_unknown_account(client):
    assert _post(client, "/accounts/active", {"account_id": "999"}).status_code == 400


def test_rules_global_with_optional_account_scope(app, client, two_accounts):
    b_id = two_accounts["b_id"]
    resp = _post(
        client,
        "/rules/new",
        {
            "name": "protect-b",
            "kind": "protected_eni",
            "pattern": "detached",
            "enabled": "on",
            "account_ref": str(b_id),
        },
    )
    assert resp.status_code == 302
    page = client.get("/rules").data.decode()
    assert "example-b" in page and "other account" in page
    # active account 1: the rule does not apply, the suggestion is allowed
    assert "protect-b" not in client.get("/suggestions").data.decode()

    _activate(client, b_id)
    page = client.get("/suggestions").data.decode()
    assert "protect-b" in page and "blocked" in page

    exported = client.get("/rules/export.yaml").data.decode()
    assert f"account_scope: {b_id}" in exported
    page = client.get("/rules/1/edit").data.decode()
    assert f'<option value="{b_id}" selected>example-b</option>' in page


@mock_aws
def test_refresh_tags_snapshot_with_active_account(app, client):
    ec2 = boto3.client("ec2", region_name="us-east-1")
    vpc_id = ec2.create_vpc(CidrBlock="10.0.0.0/16")["Vpc"]["VpcId"]
    keys_id = _add_account(
        client,
        display_name="example-keys",
        auth_mode="keys",
        access_key_id=FAKE_KEY_ID,
        secret_access_key=FAKE_SECRET,
    )
    _activate(client, keys_id)
    page = _post(client, "/refresh", follow_redirects=True).data.decode()
    assert "Refreshed example-keys (123456789012)" in page and vpc_id in page
    _activate(client, 1)
    assert vpc_id not in client.get("/").data.decode()
    with closing(_db(app)) as conn:
        refs = [r[0] for r in conn.execute("SELECT account_ref FROM snapshots")]
    assert refs == [keys_id]
    _assert_no_secrets(_log_text(app))


def test_delete_account_switches_active(app, client, two_accounts):
    _activate(client, two_accounts["b_id"])
    _post(client, f"/accounts/{two_accounts['b_id']}/delete")
    store = app.extensions["iplens"]["store"]
    assert store.active_account_id() == 1
    with closing(_db(app)) as conn:
        assert (
            conn.execute(
                "SELECT COUNT(*) FROM snapshots WHERE id=?", (two_accounts["b"].id,)
            ).fetchone()[0]
            == 0
        )
    _post(client, "/accounts/1/delete")
    page = client.get("/").data.decode()
    assert "add an AWS account" in page
    page = _post(client, "/refresh", follow_redirects=True).data.decode()
    assert "Add an AWS account in Settings first." in page


# -- history ----------------------------------------------------------------------------


def _snapshot_ids(app) -> set[int]:
    with closing(_db(app)) as conn:
        return {r[0] for r in conn.execute("SELECT id FROM snapshots")}


@pytest.fixture
def history(app, client, snapshot_builder):
    b_id = _add_account(client, display_name="example-b")
    a = [snapshot_builder(_db(app), account_ref=1).id for _ in range(3)]
    a_failed = snapshot_builder(_db(app), account_ref=1, status="failed").id
    b = [snapshot_builder(_db(app), account_ref=b_id).id for _ in range(2)]
    _activate(client, 1)
    return {"a": a, "a_failed": a_failed, "b": b, "b_id": b_id}


def test_history_lists_active_account_with_latest_protected(client, history):
    page = client.get("/").data.decode()
    table = page.split("Recent snapshots", 1)[1]
    for sid in (*history["a"][:2], history["a_failed"]):
        assert f'name="snapshot_id" value="{sid}"' in table
    # the latest successful snapshot has no checkbox (newer failed one does)
    assert f'name="snapshot_id" value="{history["a"][2]}"' not in table
    assert f'<tr title="Snapshot #{history["b"][0]}">' not in table
    assert 'Type <span class="mono">DELETE</span> to confirm' in table


def test_history_delete_requires_typed_confirmation(app, client, history):
    before = _snapshot_ids(app)
    for confirm in ("", "delete", "yes"):
        page = _post(
            client,
            "/snapshots/delete",
            {"snapshot_id": str(history["a"][0]), "confirm": confirm},
            follow_redirects=True,
        ).data.decode()
        assert "Type DELETE to confirm." in page
    assert _snapshot_ids(app) == before
    assert _post(client, "/snapshots/clear", {"confirm": "nope"}).status_code == 302
    assert _snapshot_ids(app) == before


def test_history_delete_selected_keeps_latest_and_other_accounts(app, client, history):
    a, b = history["a"], history["b"]
    resp = client.post(
        "/snapshots/delete",
        data={
            "csrf_token": _csrf(client),
            "confirm": "DELETE",
            # the latest of account 1 and a snapshot of account B are refused
            "snapshot_id": [str(a[0]), str(a[2]), str(history["a_failed"]), str(b[0])],
        },
        follow_redirects=True,
    )
    page = resp.data.decode()
    assert "Deleted 2 snapshot(s)." in page
    assert "The latest snapshot of each account is always kept." in page
    assert _snapshot_ids(app) == {a[1], a[2], *b}


def test_clear_history_this_account(app, client, history):
    page = _post(
        client, "/snapshots/clear", {"confirm": "DELETE", "scope": "account"}, follow_redirects=True
    ).data.decode()
    assert "deleted 3 snapshot(s)" in page
    assert _snapshot_ids(app) == {history["a"][2], *history["b"]}


def test_clear_history_all_accounts_keeps_latest_per_account(app, client, history):
    _post(client, "/snapshots/clear", {"confirm": "DELETE", "scope": "all"})
    assert _snapshot_ids(app) == {history["a"][2], history["b"][1]}
    # the views still have data for both accounts
    assert "No snapshot yet" not in client.get("/").data.decode()
    _activate(client, history["b_id"])
    assert "No snapshot yet" not in client.get("/").data.decode()


# -- Visual: borders, saved layout, exports -------------------------------------------------


def test_visual_border_prefs_remembered_per_account(client, two_accounts):
    page = client.get("/visual").data.decode()
    assert 'id="show-vpc" checked' in page and 'id="show-subnets" checked' in page
    assert _post(client, "/visual/prefs", {"show_vpc": "0", "show_subnets": "1"}).status_code == 204
    page = client.get("/visual").data.decode()
    assert 'id="show-vpc" >' in page and 'id="show-subnets" checked' in page

    _activate(client, two_accounts["b_id"])
    page = client.get("/visual").data.decode()
    assert 'id="show-vpc" checked' in page  # account B keeps the defaults


def test_visual_shorten_names_pref_remembered_per_account(client, two_accounts):
    page = client.get("/visual").data.decode()
    assert 'id="shorten-names" >' in page  # unchecked by default
    assert 'data-short-max="32"' in page
    assert _post(client, "/visual/prefs", {"shorten_names": "1"}).status_code == 204
    page = client.get("/visual").data.decode()
    assert 'id="shorten-names" checked' in page
    # saving only the shorten toggle leaves the border toggles alone, and vice versa
    assert 'id="show-vpc" checked' in page and 'id="show-subnets" checked' in page
    assert _post(client, "/visual/prefs", {"show_vpc": "0", "show_subnets": "1"}).status_code == 204
    page = client.get("/visual").data.decode()
    assert 'id="shorten-names" checked' in page and 'id="show-vpc" >' in page

    _activate(client, two_accounts["b_id"])
    assert 'id="shorten-names" >' in client.get("/visual").data.decode()  # B: default
    assert _post(client, "/visual/prefs", {"shorten_names": "1"}).status_code == 204
    assert _post(client, "/visual/prefs", {"shorten_names": "0"}).status_code == 204
    assert 'id="shorten-names" >' in client.get("/visual").data.decode()

    _activate(client, 1)
    assert 'id="shorten-names" checked' in client.get("/visual").data.decode()


def test_visual_layout_saved_per_account_and_vpc(client, two_accounts):
    positions = {f"res:{ENI_A}": {"x": 120.5, "y": -40}, f"subnet:{SUBNET_A}": {"x": 0, "y": 0}}
    resp = _post(client, "/visual/layout", {"vpc": VPC_A, "positions": json.dumps(positions)})
    assert resp.status_code == 204
    page = client.get(f"/visual?vpc={VPC_A}").data.decode()
    state = page.split('id="visual-positions">', 1)[1].split("</script>", 1)[0]
    assert json.loads(state) == {
        f"res:{ENI_A}": {"x": 120.5, "y": -40.0},
        f"subnet:{SUBNET_A}": {"x": 0.0, "y": 0.0},
    }

    _activate(client, two_accounts["b_id"])
    page = client.get("/visual").data.decode()
    assert 'id="visual-positions">{}</script>' in page

    _activate(client, 1)
    assert _post(client, "/visual/layout/reset", {"vpc": VPC_A}).status_code == 204
    assert 'id="visual-positions">{}</script>' in client.get("/visual").data.decode()


def _evidence_ticked(page: str) -> set[str]:
    return {
        level
        for level in ("observed", "configured", "permitted", "referenced")
        if f'name="evidence" value="{level}" checked' in page
    }


def test_extended_evidence_filter_default_and_remembered_per_account(client, two_accounts):
    url = f"/visual?vpc={VPC_A}&view=extended"
    # Default: configured + observed; permitted and referenced are opt-in.
    assert _evidence_ticked(client.get(url).data.decode()) == {"observed", "configured"}
    assert _post(client, "/visual/prefs", {"evidence": "observed,permitted"}).status_code == 204
    assert _evidence_ticked(client.get(url).data.decode()) == {"observed", "permitted"}
    # Saving other toggles keeps the evidence filter, and vice versa.
    assert _post(client, "/visual/prefs", {"shorten_names": "1"}).status_code == 204
    page = client.get(url).data.decode()
    assert _evidence_ticked(page) == {"observed", "permitted"}
    assert 'id="shorten-names" checked' in page
    assert _post(client, "/visual/prefs", {"evidence": "observed,bogus"}).status_code == 400

    _activate(client, two_accounts["b_id"])
    assert _evidence_ticked(client.get("/visual?view=extended").data.decode()) == {
        "observed",
        "configured",
    }
    assert _post(client, "/visual/prefs", {"evidence": ""}).status_code == 204  # none ticked
    assert _evidence_ticked(client.get("/visual?view=extended").data.decode()) == set()

    _activate(client, 1)
    assert _evidence_ticked(client.get(url).data.decode()) == {"observed", "permitted"}


def test_extended_view_keeps_its_own_layout(client, two_accounts):
    positions = {f"res:{ENI_A}": {"x": 10, "y": 20}}
    key = f"{VPC_A}:extended"
    resp = _post(client, "/visual/layout", {"vpc": key, "positions": json.dumps(positions)})
    assert resp.status_code == 204
    page = client.get(f"/visual?vpc={VPC_A}&view=extended").data.decode()
    assert f'data-layout-key="{key}"' in page
    assert f"res:{ENI_A}" in page.split('id="visual-positions">', 1)[1].split("</script>", 1)[0]
    page = client.get(f"/visual?vpc={VPC_A}").data.decode()  # the IP view is unaffected
    assert f'data-layout-key="{VPC_A}"' in page and 'id="visual-positions">{}</script>' in page


@pytest.mark.parametrize(
    "positions",
    ["not json", "[]", '{"a": {"x": "1", "y": 2}}', '{"a": {"x": 1e99, "y": 2}}', '{"a": 1}'],
)
def test_visual_layout_rejects_bad_positions(client, two_accounts, positions):
    resp = _post(client, "/visual/layout", {"vpc": VPC_A, "positions": positions})
    assert resp.status_code == 400


def test_visual_state_endpoints_need_csrf(client, two_accounts):
    for url in ("/visual/prefs", "/visual/layout", "/visual/layout/reset", "/visual/export.svg"):
        assert client.post(url, data={"vpc": VPC_A}).status_code == 400, url


def test_visual_page_controls_legend_and_no_cdn(client, two_accounts):
    page = client.get("/visual").data.decode()
    for marker in (
        'id="reset-layout"',
        'id="export-svg"',
        'id="export-drawio"',
        'data-export-svg="/visual/export.svg"',
        'data-export-drawio="/visual/export.drawio"',
        'id="legend-title"',
        "Arch_Amazon-EC2_48.svg",
        "Res_Elastic-Load-Balancing_Application-Load-Balancer_48.svg",
        'class="edge-swatch edge-reach"',
        'class="node-swatch node-idle"',
    ):
        assert marker in page, marker
    assert "cdn" not in page.lower() and "https://" not in page


def _export_view() -> str:
    return json.dumps(
        {
            "vpc_id": VPC_A,
            "show_vpc": True,
            "show_subnets": False,
            "nodes": [
                {
                    "id": "vpc",
                    "kind": "vpc",
                    "label": "example-vpc-a",
                    "x": 0,
                    "y": 0,
                    "w": 400,
                    "h": 300,
                    "icon": "Virtual-private-cloud-VPC_32.svg",
                },
                {
                    "id": f"subnet:{SUBNET_A}",
                    "kind": "subnet",
                    "parent": "vpc",
                    "label": "a",
                    "x": 20,
                    "y": 40,
                    "w": 300,
                    "h": 200,
                },
                {
                    "id": f"res:{ENI_A}",
                    "kind": "res",
                    "parent": f"subnet:{SUBNET_A}",
                    "label": "example\n10.0.1.10",
                    "x": 40,
                    "y": 60,
                    "w": 44,
                    "h": 44,
                    "icon": "Arch_Amazon-EC2_48.svg",
                    "idle": True,
                },
            ],
            "edges": [],
        }
    )


@pytest.mark.parametrize(
    "url, mimetype, filename, root",
    [
        ("/visual/export.svg", "image/svg+xml", f"iplens-{VPC_A}.svg", "<svg"),
        (
            "/visual/export.drawio",
            "application/vnd.jgraph.mxfile",
            f"iplens-{VPC_A}.drawio",
            "<mxfile",
        ),
    ],
)
def test_visual_export_downloads(app, client, two_accounts, url, mimetype, filename, root):
    resp = _post(client, url, {"view": _export_view()})
    assert resp.status_code == 200
    assert resp.mimetype == mimetype
    assert resp.headers["Content-Disposition"] == f"attachment; filename={filename}"
    assert root in resp.data.decode()
    assert f"exported {filename}: 3 node(s), 0 edge(s)" in _log_text(app)


def test_visual_export_rejects_bad_view(client, two_accounts):
    resp = _post(client, "/visual/export.drawio", {"view": '{"nodes": [{"kind": "x"}]}'})
    assert resp.status_code == 400 and b"invalid view" in resp.data
