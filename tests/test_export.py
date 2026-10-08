import io

from openpyxl import load_workbook

from iplens import queries
from iplens.attribution import OWNER_LABELS
from iplens.db import closing
from iplens.export import HEADERS, ips_to_xlsx

VPC = "vpc-0example0000001"
SA = "subnet-0000000a"
SB = "subnet-0000000b"


def _seed(db_path, builder):
    b = builder(db_path)
    b.vpc(VPC, "10.0.0.0/16", name="example-vpc")
    b.subnet(SA, VPC, "10.0.1.0/24", name="example-a")
    b.subnet(SB, VPC, "10.0.2.0/24", name="example-b")
    b.eni("eni-0000000001", SA, ["10.0.1.10", "10.0.1.9"], owner_ref="i-0example0001")
    b.eni("eni-0000000002", SB, ["10.0.2.5"], owner_type="lambda", owner_ref="example-fn")
    b.eni("eni-0000000003", SB, ["10.0.2.6"], owner_type="ec2", name="=HYPERLINK(1)")
    return b


def _sheet(body: bytes):
    wb = load_workbook(io.BytesIO(body))
    assert wb.sheetnames == ["IPs"]
    return wb["IPs"]


def test_xlsx_headers_layout_and_rows(db_path, snapshot_builder):
    b = _seed(db_path, snapshot_builder)
    with closing(db_path) as conn:
        rows = queries.ip_list(conn, b.id)
    ws = _sheet(ips_to_xlsx(rows, OWNER_LABELS))
    header = [c.value for c in ws[1]]
    assert header == HEADERS
    for col in (
        "IP",
        "Subnet ID",
        "Subnet name",
        "VPC ID",
        "VPC name",
        "Resource type",
        "Resource name",
        "ENI ID",
        "Status",
    ):
        assert col in header
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == f"A1:M{ws.max_row}"  # ... Primary, Tags, Terraform
    data = [dict(zip(header, (c.value for c in r), strict=True)) for r in ws.iter_rows(min_row=2)]
    assert [d["IP"] for d in data] == ["10.0.1.9", "10.0.1.10", "10.0.2.5", "10.0.2.6"]
    first = data[0]
    assert (first["Subnet ID"], first["Subnet name"], first["VPC name"]) == (
        SA,
        "example-a",
        "example-vpc",
    )
    assert (first["Resource type"], first["ENI ID"], first["Status"]) == (
        "EC2",
        "eni-0000000001",
        "in-use",
    )
    assert data[2]["Resource type"] == "Lambda" and data[2]["Resource name"] == "example-fn"


def test_xlsx_never_writes_formulas(db_path, snapshot_builder):
    b = _seed(db_path, snapshot_builder)
    with closing(db_path) as conn:
        rows = queries.ip_list(conn, b.id, queries.IpFilter(q="10.0.2.6"))
    ws = _sheet(ips_to_xlsx(rows, OWNER_LABELS))
    cell = ws.cell(row=2, column=HEADERS.index("Resource name") + 1)
    assert cell.value == "=HYPERLINK(1)" and cell.data_type == "s"


def test_xlsx_empty_has_header_only():
    ws = _sheet(ips_to_xlsx([], OWNER_LABELS))
    assert ws.max_row == 1 and [c.value for c in ws[1]] == HEADERS
