import io
from datetime import UTC, datetime

import boto3
import pytest
from botocore.exceptions import ClientError
from markupsafe import Markup
from moto import mock_aws
from openpyxl import load_workbook

from iplens import queries
from iplens.db import closing
from iplens.rules import list_rules
from iplens.web import account_label, create_app, host_allowed, local_time, utc_iso

FAKE_KEY_ID = "AKIAEXAMPLEEXAMPLE00"
FAKE_SECRET = "example/secret/value/for/tests/only/0000"
VPC = "vpc-0example0000001"
SA = "subnet-0000000a"


@pytest.fixture
def app(home):
    return create_app(home, testing=True)


@pytest.fixture
def client(app):
    c = app.test_client()
    c.get("/")
    return c


def _csrf(client) -> str:
    with client.session_transaction() as s:
        return s["csrf"]


def _post(client, url, data=None, **kw):
    data = dict(data or {})
    data["csrf_token"] = _csrf(client)
    return client.post(url, data=data, **kw)


def _db(app):
    return app.extensions["iplens"]["paths"].db_path


@pytest.fixture
def seeded(app, snapshot_builder, ips):
    b = snapshot_builder(_db(app))
    b.vpc(VPC, "10.0.0.0/16")
    b.subnet(SA, VPC, "10.0.1.0/24", name="example-private-a")
    b.eni(
        "eni-0000000001",
        SA,
        ips("10.0.1.0", 10, 3),
        owner_ref="i-0example0001",
        description="example app server",
    )
    b.eni(
        "eni-0000000002",
        SA,
        ["10.0.1.50"],
        status="available",
        owner_type="other",
        description="example detached",
    )
    b.lambda_fn("example-no-vpc", None)
    return b


def test_pages_render_without_snapshot(client):
    for url in ("/", "/ips", "/suggestions", "/rules", "/logs", "/settings", "/rules/new"):
        resp = client.get(url)
        assert resp.status_code == 200, url
    page = client.get("/").data.decode()
    assert "No data yet for <b>Account 1</b>" in page
    assert "viewing <b>Account 1</b> · no data yet" in page


def test_post_requires_csrf(client):
    assert client.post("/refresh").status_code == 400
    assert client.post("/settings", data={"csrf_token": "wrong"}).status_code == 400


def test_overview_subnet_eni_ips(client, seeded):
    page = client.get("/").data.decode()
    assert VPC in page and SA in page and "example-private-a" in page

    page = client.get(f"/subnets/{SA}").data.decode()
    assert "cell-used" in page and "cell-idle" in page and "cell-reserved" in page
    assert "/enis/eni-0000000002" in page
    assert client.get("/subnets/subnet-ffffffff").status_code == 404

    page = client.get("/enis/eni-0000000001").data.decode()
    assert "10.0.1.10" in page and "i-0example0001" in page
    assert client.get("/enis/eni-missing").status_code == 404

    page = client.get("/ips?q=detached").data.decode()
    assert "10.0.1.50" in page and "10.0.1.10" not in page
    page = client.get("/ips?owner=ec2").data.decode()
    assert "10.0.1.12" in page and "10.0.1.50" not in page


def test_settings_log_dir_change(app, client, tmp_path):
    new_dir = tmp_path / "custom-logs"
    _post(client, "/settings", {"log_dir": str(new_dir)})
    assert app.extensions["iplens"]["log_dir"] == new_dir
    assert (new_dir / "iplens.log").exists()
    page = client.get("/logs").data.decode()
    assert str(new_dir) in page and "settings saved" in page


@mock_aws
def test_test_connection_and_refresh(client):
    ec2 = boto3.client("ec2", region_name="us-east-1")
    vpc_id = ec2.create_vpc(CidrBlock="10.0.0.0/16")["Vpc"]["VpcId"]
    subnet_id = ec2.create_subnet(VpcId=vpc_id, CidrBlock="10.0.1.0/24")["Subnet"]["SubnetId"]
    ec2.create_network_interface(
        SubnetId=subnet_id, PrivateIpAddress="10.0.1.77", Description="example detached"
    )

    resp = _post(client, "/accounts/1/test", follow_redirects=True)
    assert b"Connected to account 123456789012" in resp.data

    resp = _post(client, "/refresh", follow_redirects=True)
    page = resp.data.decode()
    assert "Refreshed 123456789012 · us-east-1:" in page
    assert vpc_id in page and subnet_id in page
    # moto has no alias and no override is set: header shows the account id only
    assert "<b>123456789012</b> · us-east-1" in page

    page = client.get("/suggestions").data.decode()
    assert "Delete detached ENI" in page

    page = client.get("/logs?q=collection").data.decode()
    assert "collection finished" in page


def test_refresh_failure_is_flashed(app, client):
    def broken(_account):
        raise ValueError("example failure")

    app.extensions["iplens"]["gateway_factory"] = broken
    resp = _post(client, "/refresh", follow_redirects=True)
    assert b"Refresh failed" in resp.data
    resp = _post(client, "/accounts/1/test", follow_redirects=True)
    assert b"Connection failed" in resp.data


@pytest.mark.parametrize("exc_type", [RuntimeError, KeyError, ClientError])
def test_refresh_unexpected_error_is_generic_in_ui_and_detailed_in_log(app, client, exc_type):
    detail = "arn:aws:iam::123456789012:role/example-detail-0001"
    if exc_type is ClientError:
        exc = ClientError({"Error": {"Code": "AccessDenied", "Message": detail}}, "DescribeVpcs")
    else:
        exc = exc_type(detail)

    def broken(_settings):
        raise exc

    app.extensions["iplens"]["gateway_factory"] = broken
    resp = _post(client, "/refresh", follow_redirects=True)
    assert resp.status_code == 200
    page = resp.data.decode()
    assert "Refresh failed. See the log for details." in page
    assert detail not in page and exc_type.__name__ not in page

    log_text = (app.extensions["iplens"]["log_dir"] / "iplens.log").read_text()
    assert "refresh failed" in log_text
    assert detail in log_text and exc_type.__name__ in log_text and "Traceback" in log_text


@pytest.mark.parametrize(
    "host, port, ok",
    [
        ("localhost", None, True),
        ("127.0.0.1:8077", None, True),
        ("localhost:8077", 8077, True),
        ("127.0.0.1:8077", 8077, True),
        ("LOCALHOST:8077", 8077, True),
        ("localhost", 8077, False),  # bare name means port 80
        ("localhost:9999", 8077, False),  # wrong port
        ("evil.example", None, False),
        ("evil.example:8077", 8077, False),  # DNS rebinding: attacker name, our port
        ("127.0.0.1.evil.example:8077", 8077, False),
        ("10.0.0.5:8077", 8077, False),
        ("localhost:abc", None, False),
        ("", None, False),
    ],
)
def test_host_allowed(host, port, ok):
    assert host_allowed(host, port) is ok


def test_dns_rebinding_host_rejected(home):
    app = create_app(home, testing=True, port=8077)
    c = app.test_client()
    assert c.get("/", headers={"Host": "localhost:8077"}).status_code == 200
    assert c.get("/", headers={"Host": "127.0.0.1:8077"}).status_code == 200
    for bad in ("evil.example:8077", "evil.example", "localhost:8078"):
        resp = c.get("/settings", headers={"Host": bad})
        assert resp.status_code == 400, bad
        assert b"invalid Host header" in resp.data
    # rejected before CSRF/route handling, so even a POST never reaches /refresh
    assert c.post("/refresh", headers={"Host": "evil.example:8077"}).status_code == 400
    log_text = (app.extensions["iplens"]["log_dir"] / "iplens.log").read_text()
    assert "untrusted Host header 'evil.example:8077'" in log_text


def test_ecs_scale_down_rule_via_form(app, client):
    resp = _post(
        client,
        "/rules/new",
        {
            "name": "ecs-allow",
            "kind": "ecs_scale_down",
            "mode": "allow",
            "pattern": "",
            "enabled": "on",
        },
    )
    assert resp.status_code == 400 and b"allow mode needs a pattern" in resp.data
    _post(
        client,
        "/rules/new",
        {
            "name": "ecs-allow",
            "kind": "ecs_scale_down",
            "mode": "allow",
            "pattern": "^sandbox/",
            "enabled": "on",
        },
    )
    with closing(_db(app)) as conn:
        rule = {r.name: r for r in list_rules(conn)}["ecs-allow"]
    assert rule.params == {"mode": "allow", "pattern": "^sandbox/"}
    page = client.get(f"/rules/{rule.id}/edit").data.decode()
    assert '<option value="allow" selected>' in page
    assert (
        "Only ECS services matching “^sandbox/” may be scaled down or deleted"
        in client.get("/rules").data.decode()
    )


def test_rules_crud_and_yaml(app, client, seeded):
    resp = _post(
        client,
        "/rules/new",
        {
            "name": "lambda-in-vpc",
            "kind": "lambda_vpc_required",
            "enabled": "on",
            "description": "policy",
        },
        follow_redirects=True,
    )
    assert b"created" in resp.data
    assert b"example-no-vpc is not VPC-attached" in resp.data

    resp = _post(client, "/rules/new", {"name": "bad", "kind": "min_free_pct", "percent": "x"})
    assert resp.status_code == 400 and b"percent must be a number" in resp.data

    _post(
        client,
        "/rules/new",
        {
            "name": "keep-detached",
            "kind": "protected_eni",
            "pattern": "example detached",
            "enabled": "on",
        },
    )
    page = client.get("/suggestions").data.decode()
    assert "blocked" in page and "keep-detached" in page

    with closing(_db(app)) as conn:
        rules = {r.name: r for r in list_rules(conn)}
    rid = rules["keep-detached"].id
    assert b"example detached" in client.get(f"/rules/{rid}/edit").data
    _post(
        client,
        f"/rules/{rid}/edit",
        {"name": "keep-detached", "kind": "protected_eni", "pattern": "something-else"},
    )
    with closing(_db(app)) as conn:
        assert {r.name: r for r in list_rules(conn)}["keep-detached"].params == {
            "pattern": "something-else"
        }
        assert not {r.name: r for r in list_rules(conn)}["keep-detached"].enabled

    exported = client.get("/rules/export.yaml")
    assert exported.mimetype == "application/x-yaml"
    yaml_text = exported.data.decode()
    assert "lambda-in-vpc" in yaml_text and "keep-detached" in yaml_text

    _post(client, f"/rules/{rid}/delete")
    with closing(_db(app)) as conn:
        assert [r.name for r in list_rules(conn)] == ["lambda-in-vpc"]

    resp = _post(
        client, "/rules/import", {"yaml": yaml_text, "replace": "on"}, follow_redirects=True
    )
    assert b"Imported 2 rule(s)" in resp.data
    resp = _post(client, "/rules/import", {"yaml": "rules: ["}, follow_redirects=True)
    assert b"Import failed" in resp.data

    assert client.get("/rules/9999/edit").status_code == 404
    assert _post(client, "/rules/9999/delete").status_code == 404


def test_rules_import_file_upload(app, client):
    data = {
        "file": (
            io.BytesIO(
                b"rules:\n  - name: free-20\n    kind: min_free_pct\n    params: {percent: 20}\n"
            ),
            "rules.yaml",
        )
    }
    resp = _post(
        client, "/rules/import", data, content_type="multipart/form-data", follow_redirects=True
    )
    assert b"Imported 1 rule(s)" in resp.data


def test_navigation_tabs_and_settings_back_link(client):
    page = client.get("/").data.decode()
    tabs = ["Discovery", "IP List", "Visual", "Rules", "Suggestions", "Settings", "Logs"]
    positions = [page.index(f">{t}</a>") for t in tabs]
    assert positions == sorted(positions)
    settings = client.get("/settings").data.decode()
    assert '<a href="/" class="back-link">← Back</a>' in settings
    assert client.get("/visual").status_code == 200


def test_ip_list_page_filters(client, seeded):
    page = client.get("/ips").data.decode()
    for col in ("Subnet", "VPC", "Resource type", "Resource name / ref", "ENI", "Status"):
        assert f"<th>{col}</th>" in page
    assert "example-private-a" in page and "Export to Excel" in page
    assert page.index("10.0.1.10") < page.index("10.0.1.12") < page.index("10.0.1.50")
    page = client.get(f"/ips?vpc={VPC}&subnet={SA}&owner=other").data.decode()
    assert "10.0.1.50" in page and "10.0.1.10" not in page
    page = client.get("/ips?vpc=vpc-0missing").data.decode()
    assert "No matches" in page


def test_ip_list_export_applies_filter(client, seeded):
    resp = client.get("/ips/export.xlsx?owner=ec2&q=10.0.1.1")
    assert resp.status_code == 200
    assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert f"iplens-ips-snapshot-{seeded.id}.xlsx" in resp.headers["Content-Disposition"]
    ws = load_workbook(io.BytesIO(resp.data))["IPs"]
    header = [c.value for c in ws[1]]
    assert header[:2] == ["IP", "Subnet ID"] and "Resource type" in header
    assert [r[0].value for r in ws.iter_rows(min_row=2)] == ["10.0.1.10", "10.0.1.11", "10.0.1.12"]
    everything = load_workbook(io.BytesIO(client.get("/ips/export.xlsx").data))["IPs"]
    assert everything.max_row == 1 + 4


def test_visual_data_endpoint(client, seeded):
    assert client.get("/visual/data.json").get_json()["vpc"]["vpc_id"] == VPC
    page = client.get("/visual").data.decode()
    assert "vendor/cytoscape.min.js" in page and "/visual/data.json?vpc=" in page
    assert "vendor/dagre.min.js" in page
    assert "cdn" not in page.lower()

    data = client.get(f"/visual/data.json?vpc={VPC}").get_json()
    assert set(data) == {
        "snapshot_id",
        "vpcs",
        "vpc",
        "icons",
        "edges",
        "edge_types",
        "edges_truncated",
    }
    assert data["vpcs"] == [{"vpc_id": VPC, "name": "example-vpc"}]
    vpc = data["vpc"]
    assert set(vpc) == {"vpc_id", "name", "label_name", "label_cidrs", "cidrs", "subnets"}
    (subnet,) = vpc["subnets"]
    assert subnet["subnet_id"] == SA and subnet["cidr"] == "10.0.1.0/24"
    assert (subnet["used"], subnet["idle"], subnet["free"]) == (3, 1, 256 - 5 - 4)
    nodes = {n["eni_id"]: n for n in subnet["items"]}
    assert set(nodes) == {"eni-0000000001", "eni-0000000002"}
    web = nodes["eni-0000000001"]
    assert web["kind"] == "resource" and web["type"] == "ec2"
    assert web["ips"] == ["10.0.1.10", "10.0.1.11", "10.0.1.12"]
    assert nodes["eni-0000000002"]["type_label"] == "ENI/other"
    assert client.get("/visual/data.json?vpc=vpc-0missing").status_code == 404


def test_visual_data_without_snapshot(client):
    assert client.get("/visual/data.json").get_json() == {
        "snapshot_id": None,
        "vpcs": [],
        "vpc": None,
    }


def test_visual_assets_are_vendored(client):
    names = {*queries.TYPE_ICONS.values(), *queries.LB_ICONS.values(), queries.VPC_ICON}
    for name in sorted(names):
        resp = client.get(f"/static/icons/aws/{name}")
        assert resp.status_code == 200 and b"<svg" in resp.data, name
        resp.close()
    resp = client.get("/static/vendor/cytoscape.min.js")
    assert resp.status_code == 200 and b"Cytoscape" in resp.data
    resp.close()
    resp = client.get("/static/vendor/dagre.min.js")
    assert resp.status_code == 200 and b"graphlib" in resp.data
    resp.close()


def _seed_edges(app, snapshot_builder):
    b = snapshot_builder(_db(app))
    b.vpc(VPC, "10.0.0.0/16").subnet(SA, VPC, "10.0.1.0/24")
    b.eni(
        "eni-00000alb0a",
        SA,
        ["10.0.1.5"],
        owner_type="elb",
        owner_ref="example-alb",
        sgs=("sg-0000alb",),
    )
    b.eni(
        "eni-00000web01",
        SA,
        ["10.0.1.10"],
        owner_ref="i-0example0001",
        instance_id="i-0example0001",
        sgs=("sg-0000web",),
    )
    b.lb_target("example-alb", "example-tg", "instance", "i-0example0001", 80)
    b.sg_ref("sg-0000web", "ingress", "sg-0000alb", "tcp/80")
    return b


def test_visual_data_edges_filter(app, client, snapshot_builder):
    _seed_edges(app, snapshot_builder)
    every = client.get("/visual/data.json").get_json()
    assert sorted(e["type"] for e in every["edges"]) == ["sg", "targets"]
    only_sg = client.get("/visual/data.json?edges=&edges=sg").get_json()
    assert [e["type"] for e in only_sg["edges"]] == ["sg"]
    assert {t["type"]: t["count"] for t in only_sg["edge_types"]} == {
        "targets": 1,
        "ecs_lb": 0,
        "reach": 0,
        "sg": 1,
    }
    assert client.get("/visual/data.json?edges=").get_json()["edges"] == []


def test_visual_page_edge_toggles(app, client, snapshot_builder):
    _seed_edges(app, snapshot_builder)
    page = client.get("/visual").data.decode()
    for etype in ("targets", "ecs_lb", "reach"):
        assert f'name="edges" value="{etype}" checked' in page
    assert 'name="edges" value="sg" >' in page  # SG refs are off by default
    assert '<input type="hidden" name="edges" value="">' in page
    # the data endpoint still sends every type so ticking SG refs needs no refetch
    every = client.get("/visual/data.json").get_json()
    assert [t["selected"] for t in every["edge_types"]] == [True, True, True, True]
    assert 'id="layout"' in page and 'value="dagre"' in page
    page = client.get(f"/visual?vpc={VPC}&edges=&edges=sg").data.decode()
    assert 'name="edges" value="sg" checked' in page
    assert 'name="edges" value="targets" >' in page  # unticked


# -- account identity ----------------------------------------------------------


@pytest.mark.parametrize(
    "account_id, alias, display, expected",
    [
        ("123456789012", "example-alias", "", "example-alias (123456789012)"),
        ("123456789012", "example-alias", "Example Prod", "Example Prod (123456789012)"),
        ("123456789012", "", "", "123456789012"),
        ("123456789012", None, "", "123456789012"),
        ("", "", "", "unknown account"),
    ],
)
def test_account_label(account_id, alias, display, expected):
    assert account_label(account_id, alias, display) == expected


def test_header_shows_account_id_only_without_alias(app, client, snapshot_builder):
    snapshot_builder(_db(app), account_alias="")
    for url in ("/", "/ips", "/settings"):
        page = client.get(url).data.decode()
        assert "<b>123456789012</b> · us-east-1" in page, url
        assert "(123456789012)" not in page


def test_header_label_is_frozen_on_the_snapshot(app, client, snapshot_builder):
    snapshot_builder(_db(app), account_alias="example-alias")
    page = client.get("/rules").data.decode()
    assert "<b>example-alias (123456789012)</b> · us-east-1" in page

    # renaming the account record never relabels what was already captured
    _post(
        client,
        "/accounts/1/edit",
        {"auth_mode": "env", "region": "us-east-1", "display_name": "Example Prod"},
    )
    page = client.get("/").data.decode()
    assert "viewing <b>example-alias (123456789012)</b> · us-east-1" in page
    assert 'value="Example Prod"' in client.get("/accounts/1/edit").data.decode()

    # the next capture freezes the new display name, which beats the alias
    snapshot_builder(_db(app), account_alias="example-alias")
    page = client.get("/").data.decode()
    assert "viewing <b>Example Prod (123456789012)</b> · us-east-1" in page
    history = page.split("Recent snapshots", 1)[1]
    assert "<td>Example Prod (123456789012)</td>" in history
    assert "<td>example-alias (123456789012)</td>" in history


def test_recent_snapshots_table_shows_account_not_snapshot_id(app, client, snapshot_builder):
    snapshot_builder(_db(app), account_alias="example-alias")
    b = snapshot_builder(_db(app), "eu-west-1")
    page = client.get("/").data.decode()
    table = page.split("Recent snapshots", 1)[1]
    assert "<th>#</th>" not in table and "Taken (UTC)" not in table
    assert "<th>Account</th><th>Region</th><th>Taken</th>" in table
    assert "<td>example-alias (123456789012)</td>" in table
    assert '<td class="mono">eu-west-1</td>' in table
    # the snapshot id survives only as a tooltip
    assert f'<tr title="Snapshot #{b.id}">' in table
    assert f"<td>{b.id}</td>" not in table


# -- local time ----------------------------------------------------------------


@pytest.mark.parametrize(
    "value, expected",
    [
        ("2026-01-02T03:04:05+00:00", "2026-01-02T03:04:05Z"),
        ("2026-01-02T05:04:05+02:00", "2026-01-02T03:04:05Z"),
        (datetime(2026, 1, 2, 3, 4, 5, tzinfo=UTC), "2026-01-02T03:04:05Z"),
        ("", None),
        (None, None),
        ("not a date", None),
    ],
)
def test_utc_iso(value, expected):
    assert utc_iso(value) == expected


def test_utc_iso_naive_log_timestamp_is_server_local_time():
    # Log lines carry the server's local time without an offset.
    naive = datetime(2026, 1, 2, 3, 4, 5)
    expected = naive.astimezone().astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")
    assert utc_iso("2026-01-02 03:04:05,123") == expected


def test_local_time_markup_and_fallback():
    html = str(local_time("2026-01-02T03:04:05+00:00"))
    assert html == (
        '<time class="localtime" datetime="2026-01-02T03:04:05Z" title="2026-01-02T03:04:05Z">'
        "2026-01-02 03:04:05 UTC</time>"
    )
    # unparseable input stays a plain str, so Jinja autoescapes it
    out = local_time("<b>bogus</b>")
    assert out == "<b>bogus</b>" and not isinstance(out, Markup)


def test_timestamps_rendered_for_browser_local_time(app, client, snapshot_builder):
    snapshot_builder(_db(app), taken_at=datetime(2026, 1, 2, 3, 4, 5, tzinfo=UTC))
    page = client.get("/").data.decode()
    stamp = '<time class="localtime" datetime="2026-01-02T03:04:05Z"'
    # header, overview card and recent-snapshots table
    assert page.count(stamp) == 3
    assert "updated " + stamp in page
    assert "2026-01-02T03:04:05+00:00" not in page  # no raw stored value
    assert '<script src="/static/localtime.js" defer></script>' in page

    logs = client.get("/logs").data.decode()
    assert '<time class="localtime" datetime="' in logs


def test_localtime_script_formats_in_browser_time_zone(client):
    js = client.get("/static/localtime.js").data.decode()
    # undefined locale and no timeZone option: the browser's locale and local time zone
    assert "new Intl.DateTimeFormat(undefined," in js
    assert "timeZone:" not in js
    assert 'querySelectorAll("time.localtime[datetime]")' in js
    assert 'new Date(el.getAttribute("datetime"))' in js


def test_nav_marks_active_page(client):
    page = client.get("/ips").data.decode()
    assert '<a href="/ips" class="active" aria-current="page">IP List</a>' in page
    assert '<a href="/logs">Logs</a>' in page
