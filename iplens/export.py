"""Excel (.xlsx) export of the flat IP list."""

from __future__ import annotations

import io
from collections.abc import Iterable
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .queries import resource_type_label

XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# (header, width, row -> value)
_COLUMNS: list[tuple[str, int, Any]] = [
    ("IP", 16, lambda r, _l: r["ip"]),
    ("Subnet ID", 26, lambda r, _l: r["subnet_id"] or ""),
    ("Subnet name", 24, lambda r, _l: r["subnet_name"] or ""),
    ("VPC ID", 24, lambda r, _l: r["vpc_id"] or ""),
    ("VPC name", 22, lambda r, _l: r["vpc_name"] or ""),
    ("Resource type", 16, lambda r, labels: resource_type_label(r, labels)),
    ("Resource name", 30, lambda r, _l: r["resource_name"] or ""),
    ("Resource ref", 30, lambda r, _l: r["owner_ref"] or ""),
    ("ENI ID", 24, lambda r, _l: r["eni_id"]),
    ("Status", 12, lambda r, _l: r["status"] or ""),
    ("Primary", 9, lambda r, _l: "yes" if r["is_primary"] else "no"),
    ("Tags", 40, lambda r, _l: "; ".join(f"{k}={v}" for k, v in (r.get("tags") or {}).items())),
    (
        "Terraform",
        40,
        lambda r, _l: "; ".join(f"{m['root']}: {m['address']}" for m in r.get("tf") or []),
    ),
]
HEADERS = [c[0] for c in _COLUMNS]


def ips_to_xlsx(rows: Iterable[dict[str, Any]], labels: dict[str, str]) -> bytes:
    """One sheet: bold frozen header row, autofilter over all data, one row per IP."""
    wb = Workbook()
    ws = wb.active
    ws.title = "IPs"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([fn(r, labels) for _h, _w, fn in _COLUMNS])
        # Tag values are user-controlled: never let a leading "=" become a formula.
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    for idx, (_h, width, _fn) in enumerate(_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_COLUMNS))}{ws.max_row}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
